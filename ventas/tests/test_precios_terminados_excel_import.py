# -*- coding: utf-8 -*-
"""Tests importación Excel precios terminados."""

from decimal import Decimal
from io import BytesIO
from pathlib import Path
from unittest.mock import MagicMock, patch

import openpyxl
from django.core.exceptions import PermissionDenied
from django.test import RequestFactory, SimpleTestCase, TestCase
from django.urls import reverse

from ventas.services.precios_terminados_excel_import import (
    DATA_START_ROW,
    HEADER_ROW,
    build_preview,
    es_segunda_calidad,
    parse_workbook,
    validar_encabezados_fila15,
    variacion_pct,
    _first_token,
    _normalize_ean,
    _parse_precio_positivo,
    items_desde_preview_seleccion,
)
from ventas.tests.fixtures.precios_import_sample_builder import (
    SAMPLE_ROWS,
    build_minimal_precios_import_xlsx_bytes,
)
from ventas.views_precios_terminados import (
    _PERMISO,
    precios_terminados_import_aplicar_view,
    precios_terminados_import_preview_view,
)

FIXTURE = Path(__file__).resolve().parent / "fixtures" / "precios_import_q1_sample.xlsx"
MINIMAL_XLSX = build_minimal_precios_import_xlsx_bytes()


def _mock_user():
    user = MagicMock(is_authenticated=True)
    user.is_superuser = False
    user.tiene_permiso = lambda p: p == _PERMISO
    user.get_permisos_totales = lambda: {_PERMISO}
    return user


class ParserEncabezadosTests(SimpleTestCase):
    def test_encabezados_ok_fixture(self):
        raw = FIXTURE.read_bytes() if FIXTURE.is_file() else MINIMAL_XLSX
        parsed = parse_workbook(raw)
        self.assertTrue(parsed.ok, parsed.error_fatal)
        self.assertEqual(len(parsed.filas), len(SAMPLE_ROWS))
        self.assertEqual(parsed.hoja, "Pedido total Q1-26")
        fila_float = next(f for f in parsed.filas if f.codigo == "907947-01")
        self.assertEqual(fila_float.precio_retail, Decimal("5025.20"))

    def test_encabezado_g_mal_rechaza(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.cell(HEADER_ROW, 4, "Articulos CODIGOS")
        ws.cell(HEADER_ROW, 7, "Precio Mal")
        ws.cell(HEADER_ROW, 8, "Precio Sugerido Publico")
        ws.cell(HEADER_ROW, 14, "EAN")
        buf = BytesIO()
        wb.save(buf)
        parsed = parse_workbook(buf.getvalue())
        self.assertFalse(parsed.ok)
        self.assertIn("columna G", parsed.error_fatal)

    def test_fila_invalida_codigo(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.cell(HEADER_ROW, 4, "Articulos CODIGOS")
        ws.cell(HEADER_ROW, 7, "Precio Retail Q1")
        ws.cell(HEADER_ROW, 8, "Precio Sugerido Publico")
        ws.cell(HEADER_ROW, 14, "EAN")
        ws.cell(DATA_START_ROW, 4, "BAD")
        ws.cell(DATA_START_ROW, 7, 100)
        ws.cell(DATA_START_ROW, 8, 200)
        buf = BytesIO()
        wb.save(buf)
        parsed = parse_workbook(buf.getvalue())
        self.assertTrue(parsed.ok)
        self.assertEqual(len(parsed.filas), 0)
        self.assertTrue(parsed.errores_fila)

    def test_corte_en_totales(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.cell(HEADER_ROW, 4, "Articulos CODIGOS")
        ws.cell(HEADER_ROW, 7, "Precio Retail Q1")
        ws.cell(HEADER_ROW, 8, "Precio Sugerido Publico")
        ws.cell(HEADER_ROW, 14, "EAN")
        ws.cell(DATA_START_ROW, 4, "935478-01")
        ws.cell(DATA_START_ROW, 7, 10)
        ws.cell(DATA_START_ROW, 8, 20)
        ws.cell(DATA_START_ROW + 1, 4, "TOTALES")
        ws.cell(DATA_START_ROW + 2, 4, "999999-99")
        ws.cell(DATA_START_ROW + 2, 7, 5)
        ws.cell(DATA_START_ROW + 2, 8, 6)
        buf = BytesIO()
        wb.save(buf)
        parsed = parse_workbook(buf.getvalue())
        self.assertEqual(len(parsed.filas), 1)

    def test_ean_numerico_y_float(self):
        self.assertEqual(_normalize_ean(8720245243285), "8720245243285")
        self.assertEqual(_normalize_ean(4050375068187.0), "4050375068187")
        self.assertEqual(_normalize_ean("8720245243285"), "8720245243285")
        retail, err = _parse_precio_positivo(5025.2000000001, campo="Precio Retail", fila=16)
        self.assertIsNone(err)
        self.assertEqual(retail, Decimal("5025.20"))

    def test_minimal_bytes_corte_en_totales(self):
        parsed = parse_workbook(MINIMAL_XLSX)
        self.assertTrue(parsed.ok)
        codigos = [f.codigo for f in parsed.filas]
        self.assertEqual(codigos.count("935478-01"), 2)
        self.assertNotIn("TOTALES", codigos)

    def test_codigo_duplicado_precios_distintos(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.cell(HEADER_ROW, 4, "Articulos CODIGOS")
        ws.cell(HEADER_ROW, 7, "Precio Retail Q1")
        ws.cell(HEADER_ROW, 8, "Precio Sugerido Publico")
        ws.cell(HEADER_ROW, 14, "EAN")
        for i, (g, h) in enumerate([(10, 20), (11, 20)]):
            r = DATA_START_ROW + i
            ws.cell(r, 4, "935478-01")
            ws.cell(r, 7, g)
            ws.cell(r, 8, h)
        buf = BytesIO()
        wb.save(buf)
        parsed = parse_workbook(buf.getvalue())
        self.assertFalse(parsed.ok)
        self.assertIn("935478-01", parsed.error_fatal)


class ResolucionTests(SimpleTestCase):
    def test_first_token(self):
        self.assertEqual(_first_token("907947-01 T110"), "907947-01")
        self.assertEqual(_first_token("907947-01 T4 Puma"), "907947-01")

    def test_excluye_2da(self):
        self.assertTrue(es_segunda_calidad(codigo_modelo="907947-00", nombre="X"))
        self.assertTrue(es_segunda_calidad(codigo_modelo="907947-01", nombre="Media 2da"))

    @patch("ventas.services.precios_terminados_excel_import._cargar_articulos_terminados")
    def test_preview_mapeo_lista1_neto_lista5_final(self, mock_load):
        mock_load.return_value = [
            {
                "id_articulo": 10,
                "id_manual": "100",
                "nombre": "935478-01 T110 Puma",
                "cod_art_prov": "935478-01 T110",
                "nro_cod_barra": "8720245243285",
                "nro_cod_barra_f": "",
                "precio1_neto": 100,
                "precio1_final": 121,
                "precio5_neto": 4000,
                "precio5_final": 8000,
            }
        ]
        from ventas.services.precios_terminados_excel_import import FilaExcel, ParseExcelResult

        parse = ParseExcelResult(
            ok=True,
            filas=[
                FilaExcel(
                    fila=16,
                    codigo="935478-01",
                    precio_retail=Decimal("4358.50"),
                    precio_sugerido=Decimal("8499"),
                    ean="8720245243285",
                )
            ],
        )
        preview = build_preview("emp_test", parse, nombre_archivo="t.xlsx")
        self.assertTrue(preview["ok"])
        talle = preview["grupos"][0]["talles"][0]
        self.assertEqual(talle["precios_guardar"]["1"]["neto"], 4358.5)
        self.assertEqual(talle["precios_guardar"]["5"]["final"], 8499.0)

    def test_items_desde_preview_seleccion(self):
        preview = {
            "grupos": [
                {
                    "codigo": "935478-01",
                    "talles": [
                        {
                            "id_articulo": 10,
                            "estado": "actualiza",
                            "precios_guardar": {"1": {"neto": 1}, "5": {"final": 2}},
                        }
                    ],
                }
            ]
        }
        items = items_desde_preview_seleccion(preview, ["935478-01"])
        self.assertEqual(len(items), 1)
        self.assertEqual(items[0]["precios"]["1"]["neto"], 1)


    @patch("ventas.services.precios_terminados_excel_import._cargar_articulos_terminados")
    def test_advertencia_cod_prov_vs_nombre(self, mock_load):
        mock_load.return_value = [
            {
                "id_articulo": 21,
                "id_manual": "21",
                "nombre": "907947-02 T4 Puma",
                "cod_art_prov": "907947-01 T110",
                "nro_cod_barra": "",
                "nro_cod_barra_f": "",
                "precio1_neto": 100,
                "precio1_final": 121,
                "precio5_neto": 4000,
                "precio5_final": 8000,
            }
        ]
        from ventas.services.precios_terminados_excel_import import FilaExcel, ParseExcelResult

        parse = ParseExcelResult(
            ok=True,
            filas=[
                FilaExcel(
                    fila=16,
                    codigo="907947-01",
                    precio_retail=Decimal("200"),
                    precio_sugerido=Decimal("400"),
                    ean="",
                )
            ],
        )
        preview = build_preview("emp", parse)
        talle = preview["grupos"][0]["talles"][0]
        self.assertEqual(talle["estado"], "advertencia")
        msg = "En administraNET, el artículo #21 tiene código de proveedor 907947-01 pero el nombre dice 907947-02."
        self.assertIn(msg, talle["advertencias"])
        self.assertIn({"tipo": "codigo_inconsistente", "mensaje": msg}, talle["motivos"])
        grupo = preview["grupos"][0]
        self.assertFalse(grupo["selected_default"])


class ImportUrlsTests(SimpleTestCase):
    def test_import_urls_reverse(self):
        self.assertEqual(
            reverse("ventas:precios_terminados_import_preview"),
            "/ventas/precios-terminados/importar/preview/",
        )
        self.assertEqual(
            reverse("ventas:precios_terminados_import_aplicar"),
            "/ventas/precios-terminados/importar/aplicar/",
        )


class ImportViewsTests(SimpleTestCase):
    def setUp(self):
        self.rf = RequestFactory()

    @patch("ventas.views_precios_terminados.build_preview")
    @patch("ventas.views_precios_terminados.parse_workbook")
    def test_preview_view_requiere_archivo(self, mock_parse, mock_preview):
        req = self.rf.post("/ventas/precios-terminados/importar/preview/")
        req.session = {"user": {"base_empresa": "emp_test"}}
        req.user = _mock_user()
        resp = precios_terminados_import_preview_view(req)
        self.assertEqual(resp.status_code, 400)

    def test_preview_sin_permiso(self):
        req = self.rf.post("/ventas/precios-terminados/importar/preview/")
        req.session = {"user": {"base_empresa": "emp_test"}}
        user = MagicMock(is_authenticated=True)
        user.is_admin = lambda: False
        user.tiene_permiso = lambda p: False
        user.roles = MagicMock(all=lambda: [])
        user.cod_usuario = "vendedor"
        req.user = user
        with self.assertRaises(PermissionDenied):
            precios_terminados_import_preview_view(req)

    @patch("ventas.views_precios_terminados.guardar_lote")
    def test_aplicar_view_llama_guardar_lote(self, mock_guardar):
        mock_guardar.return_value = {"ok": True, "actualizados": 1, "errores": []}
        preview = {
            "grupos": [
                {
                    "codigo": "935478-01",
                    "talles": [
                        {
                            "id_articulo": 10,
                            "estado": "actualiza",
                            "precios_guardar": {"1": {"neto": 1}, "5": {"final": 2}},
                        }
                    ],
                }
            ]
        }
        import json as json_mod

        req = self.rf.post(
            "/ventas/precios-terminados/importar/aplicar/",
            data=json_mod.dumps(
                {"preview": preview, "codigos_seleccionados": ["935478-01"], "nombre_archivo": "t.xlsx"}
            ),
            content_type="application/json",
        )
        req.session = {"user": {"base_empresa": "emp_test", "id_usuario": 1}}
        req.user = _mock_user()
        resp = precios_terminados_import_aplicar_view(req)
        self.assertEqual(resp.status_code, 200)
        mock_guardar.assert_called()
        self.assertIn("Synap import Excel", mock_guardar.call_args.kwargs.get("tipo_modificacion", ""))


class PreviewFixtureIntegrationTests(SimpleTestCase):
    @patch("ventas.services.precios_terminados_excel_import._cargar_articulos_terminados")
    def test_parse_fixture_sin_db(self, mock_load):
        mock_load.return_value = []
        raw = FIXTURE.read_bytes() if FIXTURE.is_file() else MINIMAL_XLSX
        parsed = parse_workbook(raw)
        self.assertTrue(parsed.ok)
        preview = build_preview("emp", parsed, nombre_archivo="precios.xlsx")
        self.assertTrue(preview["ok"])
        self.assertEqual(preview["resumen"]["codigos_total"], len(SAMPLE_ROWS))


class GuardarLoteTipoModificacionTests(SimpleTestCase):
    @patch("ventas.services.precios_articulo_legacy.aplicar_cambios_articulo")
    def test_tipo_modificacion_import(self, mock_aplicar):
        from ventas.services.precios_articulo_legacy import guardar_lote

        mock_aplicar.return_value = {"ok": True, "actualizado": True}
        guardar_lote(
            "emp",
            [{"id_articulo": 1, "precios": {"1": {"neto": 10}}}],
            tipo_modificacion="Synap import Excel: test.xlsx",
        )
        mock_aplicar.assert_called_once()
        self.assertEqual(
            mock_aplicar.call_args.kwargs.get("tipo_modificacion"),
            "Synap import Excel: test.xlsx",
        )


def test_where_base_excluye_gastos():
    from ventas.services.precios_terminados import _where_base, TIPO_PRODUCTO_TERMINADO, TIPO_PRODUCTO_2DA

    for tipo in (TIPO_PRODUCTO_TERMINADO, TIPO_PRODUCTO_2DA):
        sql, _ = _where_base("a", tipo_producto=tipo)
        assert "a.tipo_art" in sql
        assert "<> 'gasto'" in sql


def _art(**kw):
    base = {
        "id_articulo": 1,
        "id_manual": "1",
        "nombre": "907964-01 T4 Puma Crew Sock Negro",
        "cod_art_prov": "907964-01 T4",
        "nro_cod_barra": "",
        "nro_cod_barra_f": "",
        "precio1_neto": 100,
        "precio1_final": 121,
        "precio5_neto": 4000,
        "precio5_final": 6080.49,
    }
    base.update(kw)
    return base


def _preview_con(arts, *, retail="110", sugerido="6200", ean=""):
    from ventas.services.precios_terminados_excel_import import FilaExcel, ParseExcelResult

    parse = ParseExcelResult(
        ok=True,
        filas=[
            FilaExcel(
                fila=20,
                codigo="907964-01",
                precio_retail=Decimal(retail),
                precio_sugerido=Decimal(sugerido),
                ean=ean,
            )
        ],
        errores_fila=["Fila 17: Precio Retail vacío."],
    )
    with patch(
        "ventas.services.precios_terminados_excel_import._cargar_articulos_terminados",
        return_value=arts,
    ):
        return build_preview("emp", parse)


class MotivosRevisionTests(SimpleTestCase):
    def test_fmt_ars(self):
        from ventas.services.precios_terminados_excel_import import fmt_ars

        self.assertEqual(fmt_ars(Decimal("6080.49")), "$6.080,49")
        self.assertEqual(fmt_ars(9799), "$9.799,00")
        self.assertEqual(fmt_ars(Decimal("1234567.5")), "$1.234.567,50")
        self.assertEqual(fmt_ars(5), "$5,00")

    def test_motivo_salto_precio_sube(self):
        from ventas.services.precios_terminados_excel_import import motivo_salto_precio

        m = motivo_salto_precio("5", Decimal("6080.49"), Decimal("9799"))
        self.assertEqual(m["tipo"], "salto_precio")
        self.assertEqual(
            m["mensaje"],
            "La lista 5 sube 61% (de $6.080,49 a $9.799,00). Supera el 30% habitual: "
            "confirmá que el precio es correcto.",
        )

    def test_motivo_salto_precio_baja_y_umbral(self):
        from ventas.services.precios_terminados_excel_import import motivo_salto_precio

        m = motivo_salto_precio("1", Decimal("1000"), Decimal("500"))
        self.assertTrue(m["mensaje"].startswith("La lista 1 baja 50% (de $1.000,00 a $500,00)."))
        self.assertIsNone(motivo_salto_precio("1", Decimal("100"), Decimal("130")))

    def test_preview_salto_genera_motivo_y_queda_tildado(self):
        preview = _preview_con([_art()], retail="110", sugerido="9799")
        g = preview["grupos"][0]
        self.assertEqual(g["estado"], "advertencia")
        self.assertEqual([m["tipo"] for m in g["motivos"]], ["salto_precio"])
        self.assertIn("La lista 5 sube 61%", g["motivos"][0]["mensaje"])
        self.assertTrue(g["selected_default"])
        self.assertEqual(g["talles"][0]["advertencias"], [g["motivos"][0]["mensaje"]])

    def test_preview_ean_no_cargado_informativo(self):
        preview = _preview_con([_art()], ean="7790001")
        g = preview["grupos"][0]
        self.assertEqual(
            g["motivos"][0]["mensaje"],
            "El EAN del Excel (7790001) no está cargado en ningún talle. El precio se aplica igual por el código.",
        )
        self.assertEqual(g["motivos"][0]["tipo"], "ean_no_cargado")
        self.assertTrue(g["selected_default"])
        self.assertEqual(g["advertencias"], [g["motivos"][0]["mensaje"]])

    def test_preview_ean_otro_articulo_destildado(self):
        otro = _art(
            id_articulo=9,
            nombre="935478-01 T110 Puma Otro",
            cod_art_prov="935478-01 T110",
            nro_cod_barra="7790001",
        )
        preview = _preview_con([_art(), otro], ean="7790001")
        g = preview["grupos"][0]
        self.assertEqual(g["motivos"][0]["tipo"], "ean_otro_articulo")
        self.assertEqual(
            g["motivos"][0]["mensaje"],
            "El EAN del Excel (7790001) pertenece a 935478-01 (935478-01 T110 Puma Otro). "
            "Revisá que el código del Excel sea el correcto.",
        )
        self.assertFalse(g["selected_default"])

    def test_descripcion_talle_y_rangos(self):
        arts = [
            _art(),
            _art(id_articulo=2, nombre="907964-01 T110 Puma Crew Sock Negro", cod_art_prov="907964-01 T110", precio1_neto=90),
        ]
        preview = _preview_con(arts)
        g = preview["grupos"][0]
        self.assertEqual(g["descripcion"], "Puma Crew Sock Negro")
        self.assertEqual(g["cantidad_talles"], 2)
        self.assertEqual([t["talle"] for t in g["talles"]], ["T4", "T110"])
        self.assertEqual(g["lista1_neto_actual"], {"min": 90.0, "max": 100.0})
        self.assertEqual(g["estado"], "actualiza")
        self.assertTrue(g["selected_default"])
        self.assertEqual(g["motivos"], [])

    def test_errores_fila_detalle_y_no_encontrado(self):
        preview = _preview_con([])
        g = preview["grupos"][0]
        self.assertEqual(g["estado"], "no_encontrado")
        self.assertEqual(g["motivos"][0]["mensaje"], "No existe en BestSox.")
        self.assertFalse(g["selected_default"])
        det = preview["errores_fila_detalle"][0]
        self.assertEqual(det["fila"], 17)
        self.assertEqual(det["motivo"], "Precio Retail vacío.")

    def test_descripcion_sin_talle(self):
        from ventas.services.precios_terminados_excel_import import descripcion_sin_talle

        self.assertEqual(descripcion_sin_talle("907964-01 T4 Puma Crew Sock Negro"), "Puma Crew Sock Negro")
        self.assertEqual(descripcion_sin_talle("907964-01 Puma T110 Sock"), "Puma Sock")
        self.assertEqual(descripcion_sin_talle("Media Lisa"), "Media Lisa")

    def test_selected_default_reglas(self):
        from ventas.services.precios_terminados_excel_import import selected_default_grupo

        self.assertTrue(selected_default_grupo("actualiza", []))
        self.assertTrue(selected_default_grupo("advertencia", [{"tipo": "salto_precio"}, {"tipo": "ean_no_cargado"}]))
        self.assertFalse(selected_default_grupo("advertencia", [{"tipo": "codigo_inconsistente"}]))
        self.assertFalse(selected_default_grupo("advertencia", [{"tipo": "ean_otro_articulo"}]))
        self.assertFalse(selected_default_grupo("no_encontrado", []))
        self.assertFalse(selected_default_grupo("sin_cambios", []))
