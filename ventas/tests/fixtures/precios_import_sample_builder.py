# -*- coding: utf-8 -*-
"""Genera un Excel mínimo para tests de importación (sin datos bancarios ni PII)."""

from __future__ import annotations

from io import BytesIO
from typing import Any, List, Tuple

import openpyxl

HEADER_ROW = 15
DATA_START_ROW = 16
COL_CODIGO = 4
COL_RETAIL = 7
COL_SUGERIDO = 8
COL_EAN = 14
SHEET_PREFERIDA = "Pedido total Q1-26"

# (codigo, retail, sugerido, ean) — EAN puede ser str, int o float
SAMPLE_ROWS: List[Tuple[str, Any, Any, Any]] = [
    ("935478-01", 4358.50, 8499, "8720245243285"),
    ("907947-01", 5025.2000000001, 9900.0, "4050375068187"),
    ("908001-02", 1200, 2400, 8720245243285),
    ("908002-03", 800.25, 1500, "7790123456789"),
    ("908003-04", 650, 1299, ""),
    ("908004-05", 900, 1800, 7790123456789.0),
    ("908005-06", 750.5, 1499, "7790123456790"),
    ("908006-07", 1100, 2200, "7790123456791"),
    ("908007-08", 990, 1980, "7790123456792"),
    # Mismo código y precios (dos filas Excel; resolución a varios talles en DB)
    ("935478-01", 4358.50, 8499, "8720245243285"),
]

FOOTER_ROW_TEXT = (
    "Documento de prueba Synap — precios ficticios. "
    "No incluye CBU, cuentas bancarias ni contactos reales."
)


def build_minimal_precios_import_xlsx_bytes() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = SHEET_PREFERIDA

    ws.cell(HEADER_ROW, COL_CODIGO, "Articulos CODIGOS")
    ws.cell(HEADER_ROW, COL_RETAIL, "Precio Retail    Q1-2026")
    ws.cell(HEADER_ROW, COL_SUGERIDO, "Precio Sugerido Publico")
    ws.cell(HEADER_ROW, COL_EAN, "EAN")

    row = DATA_START_ROW
    for codigo, retail, sugerido, ean in SAMPLE_ROWS:
        ws.cell(row, COL_CODIGO, codigo)
        ws.cell(row, COL_RETAIL, retail)
        ws.cell(row, COL_SUGERIDO, sugerido)
        if ean != "":
            ws.cell(row, COL_EAN, ean)
        row += 1

    ws.cell(row, COL_CODIGO, "TOTALES")
    ws.cell(row + 2, 1, FOOTER_ROW_TEXT)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def write_minimal_fixture(path) -> None:
    path = path.__fspath__() if hasattr(path, "__fspath__") else str(path)
    with open(path, "wb") as f:
        f.write(build_minimal_precios_import_xlsx_bytes())
