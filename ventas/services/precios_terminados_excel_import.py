# -*- coding: utf-8 -*-
"""Importación de precios terminados desde Excel (formato Pedido total Q1)."""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from typing import Any, Dict, List, Optional, Sequence, Tuple

import openpyxl
from openpyxl.utils import get_column_letter

from core.mysql_pool import mysql_cursor
from core.utils.administranet_types import (
    str_codigo_manual_articulo,
    str_or_default,
    to_decimal_or_none,
    to_int_or_none,
)
from ventas.services.precios_articulo_legacy import q2
from ventas.services.precios_terminados import TIPO_PRODUCTO_TERMINADO, _where_base

HEADER_ROW = 15
DATA_START_ROW = 16
COL_CODIGO = 4  # D
COL_RETAIL = 7  # G
COL_SUGERIDO = 8  # H
COL_EAN = 14  # N

SHEET_PREFERIDA = "Pedido total Q1-26"
TOTALES_MARK = "TOTALES"

EXPECTED_D = "articulos codigos"
EXPECTED_H = "precio sugerido publico"
EXPECTED_N = "ean"
EXPECTED_G_PREFIX = "precio retail"

CODIGO_RE = re.compile(r"^\d{6}-\d{2}$", re.IGNORECASE)
VAR_ADVERTENCIA_PCT = Decimal("30")


def _normalize_header(value: Any) -> str:
    if value is None:
        return ""
    return " ".join(str(value).strip().split()).lower()


def _normalize_codigo(value: Any) -> str:
    if value is None:
        return ""
    s = str(value).strip().upper()
    return s


def _normalize_ean(value: Any) -> str:
    if value is None or value == "":
        return ""
    if isinstance(value, (int, float)):
        if isinstance(value, float) and value.is_integer():
            return str(int(value))
        return str(value).strip()
    s = str(value).strip()
    if s.endswith(".0"):
        try:
            return str(int(float(s)))
        except ValueError:
            pass
    return s


def _parse_precio_positivo(value: Any, *, campo: str, fila: int) -> Tuple[Optional[Decimal], Optional[str]]:
    if value is None or (isinstance(value, str) and not value.strip()):
        return None, f"Fila {fila}: {campo} vacío."
    try:
        if isinstance(value, str):
            value = value.replace(",", ".")
        dec = Decimal(str(value))
    except (InvalidOperation, ValueError):
        return None, f"Fila {fila}: {campo} no numérico ({value!r})."
    if dec <= 0:
        return None, f"Fila {fila}: {campo} debe ser mayor a 0."
    return q2(dec), None


def validar_encabezados_fila15(ws) -> Tuple[bool, Optional[str]]:
    got_d = _normalize_header(ws.cell(HEADER_ROW, COL_CODIGO).value)
    got_g = _normalize_header(ws.cell(HEADER_ROW, COL_RETAIL).value)
    got_h = _normalize_header(ws.cell(HEADER_ROW, COL_SUGERIDO).value)
    got_n = _normalize_header(ws.cell(HEADER_ROW, COL_EAN).value)

    if got_d != EXPECTED_D:
        return False, (
            f"La columna D debe ser «Articulos CODIGOS», se encontró «"
            f"{ws.cell(HEADER_ROW, COL_CODIGO).value}»."
        )
    if not got_g.startswith(EXPECTED_G_PREFIX):
        return False, (
            f"La columna G debe ser «Precio Retail …», se encontró «"
            f"{ws.cell(HEADER_ROW, COL_RETAIL).value}»."
        )
    if got_h != EXPECTED_H:
        return False, (
            f"La columna H debe ser «Precio Sugerido Publico», se encontró «"
            f"{ws.cell(HEADER_ROW, COL_SUGERIDO).value}»."
        )
    if got_n != EXPECTED_N:
        return False, (
            f"La columna N debe ser «EAN», se encontró «"
            f"{ws.cell(HEADER_ROW, COL_EAN).value}»."
        )
    return True, None


def _first_token(text: Any) -> str:
    s = str_or_default(text, "").strip()
    if not s:
        return ""
    return s.split()[0].lower()


def es_segunda_calidad(*, codigo_modelo: str, nombre: str) -> bool:
    if codigo_modelo.lower().endswith("-00"):
        return True
    if "2da" in (nombre or "").lower():
        return True
    return False


def variacion_pct(anterior: Decimal, nuevo: Decimal) -> Optional[Decimal]:
    if anterior is None or nuevo is None:
        return None
    if anterior == 0:
        if nuevo == 0:
            return Decimal("0")
        return Decimal("100")
    return q2(abs((nuevo - anterior) / anterior * Decimal("100")))


def fmt_ars(value: Any) -> str:
    """Formato pesos argentino: $6.080,49."""
    try:
        dec = q2(Decimal(str(value)))
    except (InvalidOperation, ValueError, TypeError):
        return "$—"
    neg = dec < 0
    entero, _, dec_part = f"{abs(dec):.2f}".partition(".")
    grupos = []
    while len(entero) > 3:
        grupos.insert(0, entero[-3:])
        entero = entero[:-3]
    grupos.insert(0, entero)
    return f"{'-' if neg else ''}${'.'.join(grupos)},{dec_part}"


def fmt_pct_ar(value: Any) -> str:
    """Porcentaje redondeado a entero (o 1 decimal si < 10)."""
    dec = Decimal(str(value))
    if dec >= 10:
        return f"{dec.quantize(Decimal('1'))}%"
    return f"{dec.quantize(Decimal('0.1'))}%".replace(".", ",")


TALLE_RE = re.compile(r"^T\d+[A-Z]?$|^T[A-Z]{1,3}$", re.IGNORECASE)


def talle_desde_articulo(art: Dict[str, Any]) -> str:
    """Token de talle (T4, T110…) desde NombreArticulo o CodArtProv."""
    for key in ("nombre", "cod_art_prov"):
        toks = str_or_default(art.get(key), "").split()
        if len(toks) >= 2 and TALLE_RE.match(toks[1]):
            return toks[1].upper()
    toks = str_or_default(art.get("cod_art_prov"), "").split()
    return toks[1].upper() if len(toks) >= 2 else ""


def descripcion_sin_talle(nombre: Any, codigo: str = "") -> str:
    """'907964-01 T4 Puma Crew Sock Negro' -> 'Puma Crew Sock Negro'."""
    toks = str_or_default(nombre, "").split()
    if toks and (CODIGO_RE.match(toks[0]) or (codigo and toks[0].upper() == codigo.upper())):
        toks = toks[1:]
    toks = [t for t in toks if not TALLE_RE.match(t)]
    return " ".join(toks).strip()


def descripcion_grupo(arts: Sequence[Dict[str, Any]], codigo: str) -> str:
    conteo: Dict[str, int] = {}
    for a in arts:
        d = descripcion_sin_talle(a.get("nombre"), codigo)
        if d:
            conteo[d] = conteo.get(d, 0) + 1
    if not conteo:
        return ""
    return max(conteo.items(), key=lambda kv: (kv[1], -len(kv[0])))[0]


def motivo_salto_precio(lista: str, actual: Decimal, nuevo: Decimal) -> Optional[Dict[str, str]]:
    """Motivo si la variación supera VAR_ADVERTENCIA_PCT; None si no."""
    if actual is None or nuevo is None:
        return None
    var = variacion_pct(actual, nuevo) or Decimal("0")
    if var <= VAR_ADVERTENCIA_PCT:
        return None
    if actual == 0:
        msg = (
            f"La lista {lista} no tenía precio y pasa a {fmt_ars(nuevo)}. "
            "Confirmá que el precio es correcto."
        )
    else:
        verbo = "sube" if nuevo > actual else "baja"
        msg = (
            f"La lista {lista} {verbo} {fmt_pct_ar(var)} (de {fmt_ars(actual)} a {fmt_ars(nuevo)}). "
            f"Supera el {VAR_ADVERTENCIA_PCT.quantize(Decimal('1'))}% habitual: confirmá que el precio es correcto."
        )
    return {"tipo": "salto_precio", "mensaje": msg}


MOTIVOS_IDENTIFICACION = ("ean_otro_articulo", "codigo_inconsistente")


def selected_default_grupo(estado: str, motivos: Sequence[Dict[str, str]]) -> bool:
    if estado not in ("actualiza", "advertencia"):
        return False
    return not any(m.get("tipo") in MOTIVOS_IDENTIFICACION for m in motivos)


def _dedupe_motivos(motivos: Sequence[Dict[str, str]]) -> List[Dict[str, str]]:
    out, seen = [], set()
    for m in motivos:
        k = (m.get("tipo"), m.get("mensaje"))
        if k in seen:
            continue
        seen.add(k)
        out.append(m)
    return out


def errores_fila_estructurados(errores: Sequence[str]) -> List[Dict[str, Any]]:
    out = []
    for e in errores:
        m = re.match(r"^Fila (\d+):\s*(.*)$", e or "")
        out.append(
            {
                "fila": int(m.group(1)) if m else None,
                "tipo": "error_fila",
                "mensaje": e,
                "motivo": m.group(2) if m else e,
            }
        )
    return out


def _rango(vals: Sequence[float]) -> Dict[str, Any]:
    if not vals:
        return {"min": None, "max": None}
    return {"min": min(vals), "max": max(vals)}


@dataclass
class FilaExcel:
    fila: int
    codigo: str
    precio_retail: Decimal
    precio_sugerido: Decimal
    ean: str


@dataclass
class ParseExcelResult:
    ok: bool
    error_fatal: Optional[str] = None
    hoja: str = ""
    filas: List[FilaExcel] = field(default_factory=list)
    errores_fila: List[str] = field(default_factory=list)
    errores_codigo: List[str] = field(default_factory=list)


def parse_workbook(
    file_bytes: bytes,
    *,
    hoja_preferida: str = SHEET_PREFERIDA,
) -> ParseExcelResult:
    try:
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    except Exception as exc:
        return ParseExcelResult(ok=False, error_fatal=f"No se pudo leer el Excel: {exc}")

    hoja = hoja_preferida if hoja_preferida in wb.sheetnames else wb.sheetnames[0]
    ws = wb[hoja]

    ok_hdr, err_hdr = validar_encabezados_fila15(ws)
    if not ok_hdr:
        wb.close()
        return ParseExcelResult(ok=False, error_fatal=err_hdr, hoja=hoja)

    filas: List[FilaExcel] = []
    errores_fila: List[str] = []
    errores_codigo: List[str] = []
    precios_por_codigo: Dict[str, Tuple[Decimal, Decimal]] = {}

    for row_num in range(DATA_START_ROW, (ws.max_row or DATA_START_ROW) + 1):
        raw_codigo = ws.cell(row_num, COL_CODIGO).value
        codigo = _normalize_codigo(raw_codigo)
        if not codigo:
            continue
        if codigo.upper() == TOTALES_MARK:
            break

        if not CODIGO_RE.match(codigo):
            errores_fila.append(
                f"Fila {row_num}: código inválido «{raw_codigo}» (formato NNNNNN-NN)."
            )
            continue

        retail, err_r = _parse_precio_positivo(
            ws.cell(row_num, COL_RETAIL).value, campo="Precio Retail", fila=row_num
        )
        sugerido, err_s = _parse_precio_positivo(
            ws.cell(row_num, COL_SUGERIDO).value, campo="Precio Sugerido", fila=row_num
        )
        if err_r:
            errores_fila.append(err_r)
            continue
        if err_s:
            errores_fila.append(err_s)
            continue

        ean = _normalize_ean(ws.cell(row_num, COL_EAN).value)
        par = (retail, sugerido)
        prev = precios_por_codigo.get(codigo)
        if prev and prev != par:
            errores_codigo.append(
                f"Código {codigo} repetido con precios distintos (filas anteriores vs fila {row_num})."
            )
            continue
        precios_por_codigo[codigo] = par
        filas.append(
            FilaExcel(
                fila=row_num,
                codigo=codigo,
                precio_retail=retail,
                precio_sugerido=sugerido,
                ean=ean,
            )
        )

    wb.close()
    if errores_codigo:
        return ParseExcelResult(
            ok=False,
            error_fatal="; ".join(errores_codigo),
            hoja=hoja,
            errores_fila=errores_fila,
            errores_codigo=errores_codigo,
        )

    return ParseExcelResult(
        ok=True,
        hoja=hoja,
        filas=filas,
        errores_fila=errores_fila,
    )


def _nombre_tabla(cursor, nombre_lower: str) -> Optional[str]:
    cursor.execute("SHOW TABLES")
    for row in cursor.fetchall():
        nombre = (row[0] if not isinstance(row, dict) else list(row.values())[0]) or ""
        if str(nombre).strip().lower() == nombre_lower:
            return str(nombre).strip()
    return None


def _cargar_articulos_terminados(base_empresa: str) -> List[Dict[str, Any]]:
    try:
        with mysql_cursor(base_empresa, dict_cursor=True) as cursor:
            tbl = _nombre_tabla(cursor, "articulo")
            if not tbl:
                return []
            base, base_params = _where_base(tipo_producto=TIPO_PRODUCTO_TERMINADO)
            tart = tbl.replace("`", "``")
            cursor.execute(
                f"""
                SELECT a.IDArt AS id_articulo,
                       a.id_manual,
                       a.NombreArticulo AS nombre,
                       a.CodArtProv AS cod_art_prov,
                       a.NroCodBarra AS nro_cod_barra,
                       a.NroCodBarraF AS nro_cod_barra_f,
                       a.Precio1V AS precio1_neto,
                       a.Precio1VI AS precio1_final,
                       a.Precio5V AS precio5_neto,
                       a.Precio5VI AS precio5_final
                FROM `{tart}` a
                WHERE {base}
                """,
                tuple(base_params),
            )
            return [dict(r) for r in cursor.fetchall()]
    except Exception:
        return []


def _articulo_coincide_codigo(art: Dict[str, Any], codigo: str) -> bool:
    codigo_l = codigo.lower()
    tok_prov = _first_token(art.get("cod_art_prov"))
    tok_nom = _first_token(art.get("nombre"))
    return tok_prov == codigo_l or tok_nom == codigo_l


def _barcodes_art(art: Dict[str, Any]) -> List[str]:
    out: List[str] = []
    for key in ("nro_cod_barra", "nro_cod_barra_f"):
        v = _normalize_ean(art.get(key))
        if v:
            out.append(v)
    return out


def _estado_talle(
    *,
    neto_actual: Decimal,
    final_actual: Decimal,
    neto_nuevo: Decimal,
    final_nuevo: Decimal,
    advertencias: List[str],
) -> str:
    if abs(neto_actual - neto_nuevo) < Decimal("0.005") and abs(final_actual - final_nuevo) < Decimal("0.005"):
        return "sin_cambios"
    if advertencias:
        return "advertencia"
    v1 = variacion_pct(neto_actual, neto_nuevo) or Decimal("0")
    v5 = variacion_pct(final_actual, final_nuevo) or Decimal("0")
    if v1 > VAR_ADVERTENCIA_PCT or v5 > VAR_ADVERTENCIA_PCT:
        return "advertencia"
    return "actualiza"


def build_preview(
    base_empresa: str,
    parse: ParseExcelResult,
    *,
    nombre_archivo: str = "",
) -> Dict[str, Any]:
    if not parse.ok:
        return {"ok": False, "error": parse.error_fatal or "archivo_invalido"}

    articulos_db = _cargar_articulos_terminados(base_empresa)
    ean_index: Dict[str, List[Dict[str, Any]]] = {}
    for art in articulos_db:
        for bc in _barcodes_art(art):
            ean_index.setdefault(bc, []).append(art)

    grupos: List[Dict[str, Any]] = []
    excluidos: List[Dict[str, Any]] = []
    no_encontrados: List[str] = []
    total_advertencias = 0

    for fila in parse.filas:
        codigo = fila.codigo
        matches = [a for a in articulos_db if _articulo_coincide_codigo(a, codigo)]
        aplicables = []
        for art in matches:
            nombre = str_or_default(art.get("nombre"), "")
            if es_segunda_calidad(codigo_modelo=codigo, nombre=nombre):
                excluidos.append(
                    {
                        "codigo": codigo,
                        "id_articulo": art.get("id_articulo"),
                        "nombre": nombre,
                        "motivo": "2da selección",
                    }
                )
                continue
            aplicables.append(art)

        motivos_codigo: List[Dict[str, str]] = []
        if fila.ean:
            arts_ean = ean_index.get(fila.ean, [])
            if not arts_ean:
                motivos_codigo.append(
                    {
                        "tipo": "ean_no_cargado",
                        "mensaje": (
                            f"El EAN del Excel ({fila.ean}) no está cargado en ningún talle. "
                            "El precio se aplica igual por el código."
                        ),
                    }
                )
            else:
                for ae in arts_ean:
                    if not _articulo_coincide_codigo(ae, codigo):
                        otro = (
                            _first_token(ae.get("cod_art_prov")) or _first_token(ae.get("nombre"))
                        ).upper()
                        nombre_otro = str_or_default(ae.get("nombre"), "-")
                        motivos_codigo.append(
                            {
                                "tipo": "ean_otro_articulo",
                                "mensaje": (
                                    f"El EAN del Excel ({fila.ean}) pertenece a {otro} ({nombre_otro}). "
                                    "Revisá que el código del Excel sea el correcto."
                                ),
                            }
                        )
                        break
        advertencias_codigo: List[str] = [m["mensaje"] for m in motivos_codigo]

        if not aplicables:
            no_encontrados.append(codigo)
            grupos.append(
                {
                    "codigo": codigo,
                    "fila_excel": fila.fila,
                    "precio_retail": float(fila.precio_retail),
                    "precio_sugerido": float(fila.precio_sugerido),
                    "ean": fila.ean,
                    "estado": "no_encontrado",
                    "descripcion": "",
                    "cantidad_talles": 0,
                    "motivos": [{"tipo": "no_encontrado", "mensaje": "No existe en BestSox."}]
                    + motivos_codigo,
                    "advertencias": advertencias_codigo,
                    "selected_default": False,
                    "talles": [],
                }
            )
            total_advertencias += len(advertencias_codigo)
            continue

        talles_out: List[Dict[str, Any]] = []
        for art in aplicables:
            tok_p = _first_token(art.get("cod_art_prov"))
            tok_n = _first_token(art.get("nombre"))
            motivos_talle: List[Dict[str, str]] = []
            if tok_p and tok_n and tok_p != tok_n:
                motivos_talle.append(
                    {
                        "tipo": "codigo_inconsistente",
                        "mensaje": (
                            f"En administraNET, el artículo #{art.get('id_articulo')} tiene código de "
                            f"proveedor {tok_p.upper()} pero el nombre dice {tok_n.upper()}."
                        ),
                    }
                )

            neto_act = q2(Decimal(str(art.get("precio1_neto") or 0)))
            final_act = q2(Decimal(str(art.get("precio5_final") or 0)))
            neto_nuevo = fila.precio_retail
            final_nuevo = fila.precio_sugerido
            sin_cambio_precio = (
                abs(neto_act - neto_nuevo) < Decimal("0.005")
                and abs(final_act - final_nuevo) < Decimal("0.005")
            )
            if not sin_cambio_precio:
                for lista, act, nue in (("1", neto_act, neto_nuevo), ("5", final_act, final_nuevo)):
                    m = motivo_salto_precio(lista, act, nue)
                    if m:
                        motivos_talle.append(m)
            adv_talle: List[str] = list(advertencias_codigo) + [m["mensaje"] for m in motivos_talle]
            estado = _estado_talle(
                neto_actual=neto_act,
                final_actual=final_act,
                neto_nuevo=neto_nuevo,
                final_nuevo=final_nuevo,
                advertencias=adv_talle,
            )
            if adv_talle:
                total_advertencias += 1

            talles_out.append(
                {
                    "id_articulo": art.get("id_articulo"),
                    "id_manual": str_codigo_manual_articulo(art.get("id_manual")),
                    "nombre": str_or_default(art.get("nombre"), "-"),
                    "talle": talle_desde_articulo(art),
                    "lista1_neto_actual": float(neto_act),
                    "lista1_neto_nuevo": float(neto_nuevo),
                    "lista1_var_pct": float(variacion_pct(neto_act, neto_nuevo) or 0),
                    "lista5_final_actual": float(final_act),
                    "lista5_final_nuevo": float(final_nuevo),
                    "lista5_var_pct": float(variacion_pct(final_act, final_nuevo) or 0),
                    "estado": estado,
                    "advertencias": adv_talle,
                    "motivos": motivos_talle,
                    "precios_guardar": {
                        "1": {"neto": float(neto_nuevo)},
                        "5": {"final": float(final_nuevo)},
                    },
                }
            )

        estados = {t["estado"] for t in talles_out}
        if "actualiza" in estados or "advertencia" in estados:
            if "advertencia" in estados or advertencias_codigo:
                estado_grupo = "advertencia"
            else:
                estado_grupo = "actualiza"
        elif estados == {"sin_cambios"}:
            estado_grupo = "sin_cambios"
        else:
            estado_grupo = "sin_cambios"

        motivos_grupo = _dedupe_motivos(
            list(motivos_codigo)
            + [m for t in talles_out if t["estado"] != "sin_cambios" for m in t["motivos"]]
        )
        if estado_grupo == "sin_cambios":
            # Info de identificación igual se muestra, pero no hay nada a aplicar.
            motivos_grupo = [m for m in motivos_grupo if m["tipo"] != "salto_precio"]
        grupos.append(
            {
                "codigo": codigo,
                "fila_excel": fila.fila,
                "precio_retail": float(fila.precio_retail),
                "precio_sugerido": float(fila.precio_sugerido),
                "ean": fila.ean,
                "estado": estado_grupo,
                "descripcion": descripcion_grupo(aplicables, codigo),
                "cantidad_talles": len(talles_out),
                "lista1_neto_actual": _rango([t["lista1_neto_actual"] for t in talles_out]),
                "lista5_final_actual": _rango([t["lista5_final_actual"] for t in talles_out]),
                "motivos": motivos_grupo,
                "advertencias": advertencias_codigo,
                "selected_default": selected_default_grupo(estado_grupo, motivos_grupo),
                "talles": talles_out,
            }
        )

    articulos_actualizar = sum(
        1
        for g in grupos
        for t in g.get("talles") or []
        if t.get("estado") in ("actualiza", "advertencia")
    )

    return {
        "ok": True,
        "archivo": nombre_archivo,
        "hoja": parse.hoja,
        "resumen": {
            "codigos_encontrados": sum(1 for g in grupos if g.get("talles")),
            "articulos_actualizar": articulos_actualizar,
            "no_encontrados": len(no_encontrados),
            "advertencias": total_advertencias,
            "errores_fila": len(parse.errores_fila),
            "excluidos": len(excluidos),
            "codigos_total": len(parse.filas),
        },
        "grupos": grupos,
        "errores_fila": parse.errores_fila,
        "errores_fila_detalle": errores_fila_estructurados(parse.errores_fila),
        "excluidos": excluidos,
        "no_encontrados": no_encontrados,
    }


def items_desde_preview_seleccion(
    preview: Dict[str, Any],
    codigos_seleccionados: Sequence[str],
) -> List[Dict[str, Any]]:
    sel = {c.strip().upper() for c in codigos_seleccionados if c}
    items: List[Dict[str, Any]] = []
    for grupo in preview.get("grupos") or []:
        codigo = (grupo.get("codigo") or "").upper()
        if codigo not in sel:
            continue
        for talle in grupo.get("talles") or []:
            if talle.get("estado") not in ("actualiza", "advertencia", "sin_cambios"):
                continue
            if talle.get("estado") == "sin_cambios":
                continue
            aid = to_int_or_none(talle.get("id_articulo"))
            if aid is None:
                continue
            items.append(
                {
                    "id_articulo": aid,
                    "precios": talle.get("precios_guardar") or {},
                }
            )
    return items


def generar_excel_incidencias(preview: Dict[str, Any]) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Incidencias"
    ws.append(
        [
            "Tipo",
            "Código",
            "Fila Excel",
            "IDArt",
            "Detalle",
        ]
    )
    for err in preview.get("errores_fila") or []:
        ws.append(["Error fila", "", "", "", err])
    for cod in preview.get("no_encontrados") or []:
        ws.append(["No encontrado", cod, "", "", "Sin artículos terminados"])
    for ex in preview.get("excluidos") or []:
        ws.append(
            [
                "Excluido",
                ex.get("codigo"),
                "",
                ex.get("id_articulo"),
                ex.get("motivo"),
            ]
        )
    for g in preview.get("grupos") or []:
        for adv in g.get("advertencias") or []:
            ws.append(["Advertencia", g.get("codigo"), g.get("fila_excel"), "", adv])
        for t in g.get("talles") or []:
            for adv in t.get("advertencias") or []:
                ws.append(
                    [
                        "Advertencia talle",
                        g.get("codigo"),
                        g.get("fila_excel"),
                        t.get("id_articulo"),
                        adv,
                    ]
                )
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def encabezados_esperados_para_ui() -> Dict[str, str]:
    return {
        "fila": str(HEADER_ROW),
        "D": "Articulos CODIGOS",
        "G": "Precio Retail …",
        "H": "Precio Sugerido Publico",
        "N": "EAN",
        "datos_desde_fila": str(DATA_START_ROW),
    }
